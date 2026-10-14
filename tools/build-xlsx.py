#!/usr/bin/env python3
"""Сборка книги для прогонов (.xlsx) из docs/test-cases*.md.

Книга состоит из трёх листов:
  Инструкция — как проводить прогон, статусы, приоритеты, общие предусловия,
               обозначения и примеры заполнения;
  Прогон     — по строке на кейс: описание в столбцах A–I, поля для заполнения J–M;
  Сводка     — параметры прогона и итоги по статусам, разделам и приоритетам (формулы).

Использование:
  tools/build-xlsx.py                          # docs/test-cases.md → docs/test-cases.xlsx
  tools/build-xlsx.py docs/test-cases-short.md # свой файл с кейсами
  tools/build-xlsx.py docs/test-cases.md /tmp/out.xlsx

Нужен openpyxl (pip install openpyxl). После сборки книгу стоит пересчитать
в LibreOffice, иначе формулы «Сводки» будут без сохранённых значений:
  soffice --headless --convert-to xlsx --outdir <папка> <файл>
"""
import html
import math
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

FONT = 'Arial'
SIZE = 10
LAST_ROW = 500  # диапазон формул «Сводки» с запасом на дописанные кейсы

STATUSES = ['Не выполнялся', 'Пройден', 'Не пройден', 'Заблокирован', 'Пропущен']
PRIORITIES = ['Высокий', 'Средний', 'Низкий']

COLUMNS = [
    # (заголовок, ширина, поле кейса, столбец для заполнения)
    ('ID', 12, 'id', False),
    ('Название', 28, 'title', False),
    ('Раздел', 15, 'section', False),
    ('Приоритет', 11, 'priority', False),
    ('Требование', 24, 'req', False),
    ('Предусловия', 36, 'pre', False),
    ('Шаги', 46, 'steps', False),
    ('Ожидаемый результат', 46, 'expected', False),
    ('Постусловия', 34, 'post', False),
    ('Статус', 15, None, True),
    ('Фактический результат', 36, None, True),
    ('Баг-репорт', 14, None, True),
    ('Комментарий', 28, None, True),
]

# --------------------------------------------------------------- разбор Markdown

def parse(path):
    """Markdown с кейсами → (разделы, кейсы, вводные блоки)."""
    sections, cases = [], []
    meta = {'title': '', 'intro': [], 'preconditions': [], 'notation': []}
    section = case = None
    in_pre = in_intro = in_common = in_notation = False

    for line in open(path, encoding='utf-8').read().split('\n'):
        if m := re.match(r'^# (.+)$', line):
            meta['title'], in_intro = m.group(1), True
            continue
        if in_intro:
            if line.strip():
                meta['intro'].append(line.strip())
            elif meta['intro']:
                in_intro = False
            continue
        if line.startswith('## Общие предусловия'):
            in_common, in_notation = True, False
            continue
        if line.startswith('### Обозначения'):
            in_notation, in_common = True, False
            continue
        if line.startswith('## ') or line.startswith('### '):
            in_common = in_notation = False
        if (in_common or in_notation) and line.startswith('- '):
            item = line[2:]
            if m := re.match(r'^\*\*(.+?)\.?\*\*\s*[—-]?\s*(.+)$', item):
                label, text = m.group(1), m.group(2)
            else:
                label, text = '', item
            (meta['preconditions'] if in_common else meta['notation']).append((label, text))
            continue

        if m := re.match(r'^## \d+\. (.+)$', line):
            section = m.group(1)
            sections.append(section)
            case, in_pre = None, False
        elif m := re.match(r'^### (TC-[A-Z]+-\d+)\. (.+)$', line):
            case = {'id': m.group(1), 'title': m.group(2), 'section': section,
                    'pre': [], 'steps': [], 'notes': [], 'post': ''}
            cases.append(case)
            in_pre = False
        elif case is None:
            continue
        elif line.startswith('**Приоритет:**'):
            m = re.match(r'^\*\*Приоритет:\*\* (\S+) · \*\*Требование:\*\* (.+)$', line)
            case['priority'] = m.group(1).capitalize()
            case['req'] = m.group(2)
        elif line.startswith('**Предусловия:**'):
            case['pre'].append(line.removeprefix('**Предусловия:**').strip())
            in_pre = True
        elif in_pre and re.match(r'^\d+\. ', line):
            case['pre'].append(line)
        elif m := re.match(r'^\| (\d+) \| (.+) \| (.+) \|$', line):
            case['steps'].append((m.group(2).strip(), m.group(3).strip()))
            in_pre = False
        elif line.startswith('**Примечание:**'):
            case['notes'].append(line.removeprefix('**Примечание:**').strip())
            in_pre = False
        elif line.startswith('**Постусловия:**'):
            case['post'] = line.removeprefix('**Постусловия:**').strip()
            in_pre = False
        elif not line.strip():
            in_pre = False

    for c in cases:
        assert c.get('priority') in PRIORITIES, c['id']
        assert c['steps'] and c['post'] and c['pre'], c['id']
    assert meta['title'] and meta['intro'] and meta['preconditions'], 'не разобрана шапка файла'
    meta['intro'] = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', ' '.join(meta['intro']))
    return sections, cases, meta

# --------------------------------------------------------------- оформление

thin = Side(style='thin', color='BFBFBF')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(wrap_text=True, vertical='top', horizontal='center')

HEAD_FILL = PatternFill('solid', fgColor='1F3864')
INPUT_HEAD_FILL = PatternFill('solid', fgColor='BF8F00')
INPUT_FILL = PatternFill('solid', fgColor='FFF2CC')
TOTAL_FILL = PatternFill('solid', fgColor='D9E1F2')

PLAIN = InlineFont(rFont=FONT, sz=SIZE)
BOLD = InlineFont(rFont=FONT, sz=SIZE, b=True)


def font(**kw):
    return Font(name=FONT, size=kw.pop('size', SIZE), **kw)


def rich(text):
    """`код` из Markdown → жирный фрагмент."""
    parts = text.split('`')
    if len(parts) == 1:
        return text
    return CellRichText(*[TextBlock(BOLD if i % 2 else PLAIN, p) for i, p in enumerate(parts) if p])


def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text).replace('`', '').replace('**', '')


def set_width(ws, letter, width):
    """Задаёт ширину столбца.

    LibreOffice при пересохранении теряет ширину столбца, если она в точности
    совпадает с шириной соседнего слева, поэтому в таком случае добавляем
    незаметную глазу добавку.
    """
    prev = ws.column_dimensions.get(chr(ord(letter) - 1)) if letter > 'A' else None
    if prev is not None and prev.width == width:
        width += 0.1
    ws.column_dimensions[letter].width = width


def fit_height(ws, row, texts_by_col):
    """Высота строки по самой длинной ячейке: переносы считаем сами."""
    lines = 1
    for col, text in texts_by_col.items():
        per_line = max(int(ws.column_dimensions[col].width * 1.05), 1)
        lines = max(lines, sum(max(1, math.ceil(len(par) / per_line)) for par in plain(text).split('\n')))
    ws.row_dimensions[row].height = max(15, lines * 12.75 + 5)


def print_fit(ws, landscape=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if landscape:
        ws.page_setup.orientation = 'landscape'

# --------------------------------------------------------------- лист «Прогон»

def case_fields(c):
    steps = '\n'.join(f'{n}. {s}' for n, (s, _) in enumerate(c['steps'], start=1))
    expected = '\n'.join(f'{n}. {e}' for n, (_, e) in enumerate(c['steps'], start=1))
    if c['notes']:
        expected += '\n' + '\n'.join(f'Примечание: {n}' for n in c['notes'])
    return {'id': c['id'], 'title': c['title'], 'section': c['section'], 'priority': c['priority'],
            'req': c['req'], 'pre': '\n'.join(c['pre']), 'steps': steps, 'expected': expected,
            'post': c['post'][:1].upper() + c['post'][1:]}


def build_run(wb, cases):
    ws = wb.active
    ws.title = 'Прогон'
    for i, (title, width, _, is_input) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=title)
        cell.font = font(bold=True, color='FFFFFF')
        cell.fill = INPUT_HEAD_FILL if is_input else HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = BORDER
        set_width(ws, get_column_letter(i), width)
    ws.row_dimensions[1].height = 30

    for r, case in enumerate(cases, start=2):
        fields = case_fields(case)
        texts = {}
        for i, (title, _, key, is_input) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=r, column=i)
            cell.border = BORDER
            cell.font = font()
            cell.alignment = WRAP
            if is_input:
                cell.fill = INPUT_FILL
                if title == 'Статус':
                    cell.value = STATUSES[0]
                    cell.alignment = CENTER
                continue
            text = fields[key]
            texts[get_column_letter(i)] = text
            cell.value = text if key in ('id', 'section', 'priority') else rich(text)
            if key == 'id':
                cell.font = font(bold=True)
            if key == 'priority':
                cell.alignment = CENTER
        fit_height(ws, r, texts)

    status_col = get_column_letter(1 + [c[0] for c in COLUMNS].index('Статус'))
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Статус', error='Выберите статус из списка',
                        showInputMessage=True, promptTitle='Статус', prompt='Выберите статус кейса')
    ws.add_data_validation(dv)
    dv.add(f'{status_col}2:{status_col}{LAST_ROW}')
    for status, (bg, fg) in {'Пройден': ('C6EFCE', '006100'), 'Не пройден': ('FFC7CE', '9C0006'),
                             'Заблокирован': ('F8CBAD', '843C0C'), 'Пропущен': ('D9D9D9', '404040')}.items():
        ws.conditional_formatting.add(
            f'{status_col}2:{status_col}{LAST_ROW}',
            FormulaRule(formula=[f'${status_col}2="{status}"'],
                        fill=PatternFill('solid', fgColor=bg, bgColor=bg),
                        font=Font(name=FONT, color=fg, bold=True)))

    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{len(cases) + 1}'
    ws.print_title_rows = '1:1'
    print_fit(ws, landscape=True)
    return ws

# --------------------------------------------------------------- лист «Сводка»

PARAMS = [
    ('Версия сборки', 'из подвала страницы калькулятора, например 1.0.0'),
    ('Стенд', 'https://calc.simdyanov.ru, http://localhost:5173 или http://localhost:8080'),
    ('Браузер и ОС', 'например, Chrome 140, macOS 26'),
    ('Тестировщик', ''),
    ('Дата начала', 'ДД.ММ.ГГГГ'),
    ('Дата окончания', 'ДД.ММ.ГГГГ'),
]


def build_summary(wb, sections):
    ws = wb.create_sheet('Сводка')
    set_width(ws, 'A', 34)
    for col in 'BCDEFG':
        set_width(ws, col, 15)
    set_width(ws, 'H', 16)

    ws['A1'] = 'Прогон тестов: калькулятор, JS-версия'
    ws['A1'].font = font(bold=True, size=14)

    def heading(cell, text):
        ws[cell] = text
        ws[cell].font = font(bold=True, size=12)

    heading('A3', 'Параметры прогона')
    for i, (label, hint) in enumerate(PARAMS, start=4):
        ws[f'A{i}'] = label
        ws[f'A{i}'].font = font(bold=True)
        ws[f'A{i}'].border = BORDER
        ws.merge_cells(f'B{i}:D{i}')
        ws[f'B{i}'].fill = INPUT_FILL
        ws[f'B{i}'].font = font()
        for col in 'BCD':
            ws[f'{col}{i}'].border = BORDER
        if hint:
            ws[f'E{i}'] = hint
            ws[f'E{i}'].font = font(italic=True, color='7F7F7F')

    status_range = f'Прогон!$J$2:$J${LAST_ROW}'
    heading('A11', 'Итоги')
    totals = [
        ('Всего кейсов', f'=COUNTA(Прогон!$A$2:$A${LAST_ROW})', None),
        ('Пройден', f'=COUNTIF({status_range},"Пройден")', None),
        ('Не пройден', f'=COUNTIF({status_range},"Не пройден")', None),
        ('Заблокирован', f'=COUNTIF({status_range},"Заблокирован")', None),
        ('Пропущен', f'=COUNTIF({status_range},"Пропущен")', None),
        ('Не выполнялся', '=B12-SUM(B13:B16)', 'включая кейсы с пустым статусом'),
        ('Выполнено', '=IF(B12=0,0,(B13+B14)/B12)', 'доля кейсов со статусом «Пройден» или «Не пройден»'),
        ('Пройдено из выполненных', '=IF(B13+B14=0,"—",B13/(B13+B14))', '«—», пока ни один кейс не выполнен'),
    ]
    for i, (label, formula, hint) in enumerate(totals, start=12):
        ws[f'A{i}'] = label
        ws[f'A{i}'].font = font(bold=(i == 12))
        ws[f'A{i}'].border = BORDER
        ws[f'B{i}'] = formula
        ws[f'B{i}'].font = font(bold=(i == 12))
        ws[f'B{i}'].border = BORDER
        ws[f'B{i}'].alignment = Alignment(horizontal='right')
        if i >= 18:
            ws[f'B{i}'].number_format = '0%'
        if hint:
            ws[f'C{i}'] = hint
            ws[f'C{i}'].font = font(italic=True, color='7F7F7F')

    def breakdown(start_row, title, first_label, labels, column):
        heading(f'A{start_row}', title)
        head = start_row + 1
        crit = f'Прогон!${column}$2:${column}${LAST_ROW}'
        for j, text in enumerate([first_label, 'Всего'] + STATUSES[1:] + [STATUSES[0]], start=1):
            cell = ws.cell(row=head, column=j, value=text)
            cell.font = font(bold=True, color='FFFFFF')
            cell.fill = HEAD_FILL
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='left' if j == 1 else 'center',
                                       vertical='center', wrap_text=True)
        ws.row_dimensions[head].height = 28
        for k, label in enumerate(labels):
            r = head + 1 + k
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=f'=COUNTIF({crit},$A{r})')
            for j in range(3, 7):
                col = get_column_letter(j)
                ws.cell(row=r, column=j, value=f'=COUNTIFS({crit},$A{r},{status_range},{col}${head})')
            ws.cell(row=r, column=7, value=f'=B{r}-SUM(C{r}:F{r})')
        total = head + 1 + len(labels)
        ws.cell(row=total, column=1, value='Итого')
        for j in range(2, 8):
            col = get_column_letter(j)
            ws.cell(row=total, column=j, value=f'=SUM({col}{head + 1}:{col}{total - 1})')
        for r in range(head + 1, total + 1):
            for j in range(1, 8):
                cell = ws.cell(row=r, column=j)
                cell.border = BORDER
                cell.font = font(bold=(r == total))
                if j > 1:
                    cell.alignment = Alignment(horizontal='right')
                if r == total:
                    cell.fill = TOTAL_FILL
        return total

    end = breakdown(21, 'По разделам', 'Раздел', sections, 'C')
    breakdown(end + 2, 'По приоритетам', 'Приоритет', PRIORITIES, 'D')
    print_fit(ws)
    return ws

# --------------------------------------------------------------- лист «Инструкция»

HOWTO = [
    ('1. Параметры', 'На листе «Сводка» заполните жёлтые ячейки: версию сборки из подвала страницы, стенд, '
                     'браузер и ОС, своё имя, даты прогона.'),
    ('2. Выполнение', 'На листе «Прогон» выполняйте кейсы по порядку или отфильтруйте их по разделу и приоритету '
                      '(стрелки в заголовках). Перед каждым кейсом верните калькулятор в исходное состояние (ОП-1).'),
    ('3. Результаты', 'Заполняйте только жёлтые столбцы J–M: «Статус» (выбор из списка), «Фактический результат», '
                      '«Баг-репорт», «Комментарий». Столбцы A–I — описание кейса, их не меняйте.'),
    ('4. Итоги', 'Лист «Сводка» считает итоги сам: по статусам, разделам и приоритетам.'),
]

STATUS_HELP = [
    ('Не выполнялся', 'Кейс ещё не проверяли. Стоит по умолчанию.'),
    ('Пройден', 'На всех шагах фактический результат совпал с ожидаемым.'),
    ('Не пройден', 'Хотя бы на одном шаге результат отличается. В «Фактический результат» запишите номер шага '
                   'и что увидели, в «Баг-репорт» — номер или ссылку на заведённый дефект.'),
    ('Заблокирован', 'Кейс нельзя выполнить: не получается выполнить предусловие или мешает другой дефект. '
                     'Причину запишите в «Комментарий».'),
    ('Пропущен', 'Кейс сознательно не выполняли в этом прогоне, например он не входит в объём смоук-теста. '
                 'Причину запишите в «Комментарий».'),
]

PRIORITY_HELP = [
    ('Высокий', 'Основная функция: без неё калькулятором нельзя пользоваться.'),
    ('Средний', 'Важная, но не основная функция.'),
    ('Низкий', 'Граничные и редкие случаи.'),
]

EXAMPLE_PASSED = [('ID', 'TC-ARI-01'), ('Статус', 'Пройден'),
                  ('Фактический результат', 'Оставьте пустым: он совпал с ожидаемым.')]
EXAMPLE_FAILED = [('ID', 'TC-INP-02'), ('Статус', 'Не пройден'),
                  ('Фактический результат', 'Шаг 2: на экране 05, ожидалось 5'),
                  ('Баг-репорт', 'BUG-17'),
                  ('Комментарий', 'Воспроизводится в Firefox, в Chrome — нет')]


def build_guide(wb, meta, md_name, cases):
    ws = wb.create_sheet('Инструкция', 0)
    set_width(ws, 'A', 24)
    set_width(ws, 'B', 100)
    ws.sheet_view.showGridLines = False

    ws['A1'] = meta['title']
    ws['A1'].font = font(bold=True, size=14)
    ws['A2'] = f'Книга для прогонов. {meta["intro"]} Тот же набор в Markdown — docs/{md_name}.'
    ws['A2'].font = font(color='595959')

    row = 4

    def block(title, rows, header=None):
        nonlocal row
        ws[f'A{row}'] = title
        ws[f'A{row}'].font = font(bold=True, size=12)
        row += 1
        if header:
            for j, text in enumerate(header, start=1):
                cell = ws.cell(row=row, column=j, value=text)
                cell.font = font(bold=True, color='FFFFFF')
                cell.fill = HEAD_FILL
                cell.border = BORDER
            row += 1
        for label, text in rows:
            a = ws.cell(row=row, column=1, value=label)
            b = ws.cell(row=row, column=2, value=rich(text))
            for cell in (a, b):
                cell.alignment = WRAP
                cell.border = BORDER
                cell.font = font()
            a.font = font(bold=True)
            fit_height(ws, row, {'A': label, 'B': text})
            row += 1
        row += 1

    example_id = next((c['id'] for c in cases if c['priority'] == 'Высокий'), cases[0]['id'])
    block('Как провести прогон', HOWTO)
    block('Статусы', STATUS_HELP, header=['Статус', 'Когда ставить'])
    block('Приоритеты', PRIORITY_HELP)
    block('Общие предусловия', [(lbl, txt) for lbl, txt in meta['preconditions']])
    block('Обозначения', [(lbl or '—', txt[:1].upper() + txt[1:]) for lbl, txt in meta['notation']]
          + [('Шаги и результаты', 'Номер строки в «Ожидаемом результате» совпадает с номером шага в «Шагах».')])
    block('Пример: кейс пройден (условные данные)', [(k, v if k != 'ID' else example_id) for k, v in EXAMPLE_PASSED])
    block('Пример: кейс не пройден (условные данные, не результат реального прогона)', EXAMPLE_FAILED)
    print_fit(ws)
    return ws


def main():
    md = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'docs', 'test-cases.md')
    out = sys.argv[2] if len(sys.argv) > 2 else md.removesuffix('.md') + '.xlsx'
    sections, cases, meta = parse(md)

    wb = Workbook()
    build_run(wb, cases)
    build_summary(wb, sections)
    build_guide(wb, meta, os.path.basename(md), cases)
    wb.active = 0
    wb.properties.title = meta['title']
    wb.save(out)
    print(f'{out}: {len(cases)} кейсов, разделов {len(sections)}: {", ".join(sections)}')


if __name__ == '__main__':
    sys.exit(main())
